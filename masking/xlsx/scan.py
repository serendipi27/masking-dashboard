"""XLSX 유출 스캔과 구조 검증. 결과에는 위치와 종류만 담고 실제값은 절대 출력하지 않는다."""
import re
import zipfile

import openpyxl

from ..core.leak_common import digits as _digits


def scan_xlsx(path, tokens):
    hits = []
    wb = openpyxl.load_workbook(str(path))
    ws = wb.worksheets[0]
    numeric_tokens = [(k, t) for k, t in tokens if len(_digits(t)) >= 6 and _digits(t) == re.sub(r'[\s-]', '', t)]
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or not v.strip():
                continue
            low, dg = v.lower(), _digits(v)
            for kind, t in tokens:
                if t.lower() in low:
                    hits.append((f'R{cell.row}C{cell.column}', kind))
            for kind, t in numeric_tokens:
                if _digits(t) in dg and t.lower() not in low:
                    hits.append((f'R{cell.row}C{cell.column}', kind + '-digits'))
    with zipfile.ZipFile(path) as z:                 # xlsx는 zip 안 xml에 문자열이 평문으로 들어감
        # 실제 셀 데이터가 들어가는 부분만 검사함. styles.xml 등에는 색상 코드 같은 16진수 문구가 있어
        # 숫자 위주 토큰과 우연히 겹치는 오탐이 날 수 있음(예: 검정색 코드 "00000000").
        data_parts = [n for n in z.namelist() if n.startswith('xl/worksheets/') or n == 'xl/sharedStrings.xml']
        raw = b''.join(z.read(n) for n in data_parts)
    for kind, t in tokens:
        if t.encode('utf-8') in raw:
            hits.append(('raw-xml', kind))
    return sorted(set(hits))


def _shape(v):
    return re.sub(r'[가-힣]', '가', re.sub(r'[A-Za-z]', 'a', re.sub(r'\d', '9', str(v))))


def verify_structure(orig_path, new_path):
    """열 형식·데이터 형 유지 여부 검증."""
    problems = []
    wo, wn = openpyxl.load_workbook(str(orig_path)), openpyxl.load_workbook(str(new_path))
    o, n = wo.worksheets[0], wn.worksheets[0]
    if (o.max_row, o.max_column) != (n.max_row, n.max_column):
        return [f'시트 크기 변경 {o.max_row}x{o.max_column} -> {n.max_row}x{n.max_column}']
    shape_cols = {1: '승인번호', 4: '사업자번호', 9: '사업자번호'}
    same_cols = {0: '작성일자', 2: '발급일자', 3: '전송일자', 25: '품목일자', 17: '분류', 18: '종류', 19: '발급유형',
                21: '영수/청구', 28: '품목수량'}
    for r in range(o.max_row):
        for c in range(o.max_column):
            a, b = o.cell(row=r + 1, column=c + 1), n.cell(row=r + 1, column=c + 1)
            if type(a.value) is not type(b.value):
                problems.append(f'R{r + 1}C{c + 1} 데이터 형 변경 {type(a.value).__name__}->{type(b.value).__name__}')
            elif a.number_format != b.number_format:
                problems.append(f'R{r + 1}C{c + 1} 서식 변경')
            if r >= 6 and isinstance(a.value, str) and a.value.strip():
                if c in shape_cols and _shape(a.value) != _shape(b.value):
                    problems.append(f'R{r + 1}C{c + 1} {shape_cols[c]} 형식 변경')
                if c in same_cols and a.value != b.value:
                    problems.append(f'R{r + 1}C{c + 1} {same_cols[c]} 값이 바뀜')
        if r >= 6 and isinstance(o.cell(row=r + 1, column=16).value, (int, float)):
            h = n.cell(row=r + 1, column=15).value
            s = n.cell(row=r + 1, column=16).value
            t = n.cell(row=r + 1, column=17).value
            oh, os_, ot = (o.cell(row=r + 1, column=c).value for c in (15, 16, 17))
            if oh == (os_ or 0) + (ot or 0) and h != (s or 0) + (t or 0):
                problems.append(f'R{r + 1} 합계 != 공급가액+세액')
    return problems
