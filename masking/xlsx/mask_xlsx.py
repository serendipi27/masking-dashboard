"""XLSX(전자세금계산서 목록, 신형식) 마스킹. openpyxl은 셀 값만 바꾸면 서식을 그대로 보존하므로
xlrd/xlutils처럼 서식을 수동으로 복사할 필요가 없다. 값 치환 규칙(회사명·사업자번호·금액 등)은
xls용 XlsMasker와 동일해 그대로 물려받아 재사용하고, 셀 입출력만 openpyxl에 맞게 새로 짠다.

열 구조가 xls와 같다는 가정이며(같은 홈택스 내려받기, 확장자만 다름), 다르면 xls/rules.py를 갱신해야 한다.
"""
import json
from collections import Counter

import openpyxl

from ..xls.mask_xls import XlsMasker, tax_of
from ..xls.rules import AMOUNT_COLS, DATA_START_ROW, HEADER_ROW, KEEP_COLS, RECEIVER_IDENTITY_COLS, TEXT_RULES
from ..core.amounts import scale_amount


class XlsxMasker(XlsMasker):
    """값 치환 로직은 XlsMasker를 그대로 물려받고, 시트 순회만 openpyxl 방식으로 다시 구현."""

    @staticmethod
    def check_headers(ws):
        def val(r, c):
            return ws.cell(row=r + 1, column=c + 1).value

        for k, (name, _) in TEXT_RULES.items():
            if str(val(HEADER_ROW, k) or '').strip() != name:
                raise ValueError(f'열 {k} 헤더 불일치: 기대 "{name}"')
        for k, name in {**AMOUNT_COLS, **KEEP_COLS}.items():
            if str(val(HEADER_ROW, k) or '').strip() != name:
                raise ValueError(f'열 {k} 헤더 불일치: 기대 "{name}"')

    def run(self, ws, edit=True, stats=None):
        """edit=False면 수집 단계(값 변환 없이 실제값만 note), edit=True면 실제 마스킹."""
        self.check_headers(ws)
        factor = self.s.amount_factor() if edit else None
        changed = Counter()
        nrows = ws.max_row
        tot_new, tot_old = [0, 0, 0], [0, 0, 0]
        for r in range(DATA_START_ROW, nrows):
            for k, (_, rule) in TEXT_RULES.items():
                cell = ws.cell(row=r + 1, column=k + 1)
                v = cell.value
                if v is None or not str(v).strip():
                    continue
                v = str(v)
                if stats is not None:
                    stats.setdefault(k, set()).add(v)
                real = v
                if k in RECEIVER_IDENTITY_COLS:
                    real = self.receiver_canonical.setdefault(k, real)
                new = getattr(self, rule)(real)
                if edit and new != v:
                    cell.value = new
                    changed[k] += 1
            if edit:
                self._amounts_xlsx(ws, r, factor, tot_new, tot_old, changed)
        self._top_block_xlsx(ws, edit, tot_new, tot_old, changed)
        return changed

    def _num(self, ws, r, k):
        v = ws.cell(row=r + 1, column=k + 1).value
        return float(v) if isinstance(v, (int, float)) else None

    def _amounts_xlsx(self, ws, r, k, tot_new, tot_old, changed):
        h, s, t, si, ti = (self._num(ws, r, c) for c in (14, 15, 16, 30, 31))
        new = {}
        if s is not None:
            new[15] = scale_amount(s, k)
            if t is not None:
                new[16] = tax_of(new[15], s, t)
            if h is not None:
                new[14] = new[15] + new.get(16, 0) if h == s + (t or 0) else scale_amount(h, k)
            if si is not None:
                new[30] = new[15] if si == s else scale_amount(si, k)
            if ti is not None:
                new[31] = new.get(16, 0) if ti == t else tax_of(new.get(30, new[15]), si or s, ti)
        for col, val in new.items():
            ws.cell(row=r + 1, column=col + 1).value = float(val)
            changed[col] += 1
        for i, (col, old) in enumerate(((14, h), (15, s), (16, t))):
            tot_old[i] += old or 0
            tot_new[i] += new.get(col, old or 0)
        cell = ws.cell(row=r + 1, column=30)                 # 품목단가: '108,000' 형태의 문자열일 수 있음
        v = cell.value
        if isinstance(v, str) and v.strip().replace(',', '').isdigit():
            cell.value = f"{scale_amount(int(v.replace(',', '')), k):,}"
            changed[29] += 1

    def _top_block_xlsx(self, ws, edit, tot_new, tot_old, changed):
        labels = {'사업자등록번호': 'biz', '상호': 'company', '대표자명': 'person'}
        totals = {'총합계금액': 0, '총공급가액': 1, '총세액': 2}
        for r in range(DATA_START_ROW - 2):
            for k in range(ws.max_column - 1):
                v = ws.cell(row=r + 1, column=k + 1).value
                if not isinstance(v, str):
                    continue
                label = v.replace(' ', '')
                vc = ws.cell(row=r + 1, column=k + 2)
                if label in labels and isinstance(vc.value, str) and vc.value.strip():
                    new = getattr(self, labels[label])(vc.value)
                    if edit and new != vc.value:
                        vc.value = new
                        changed['top'] += 1
                elif label in totals and edit and vc.value is not None:
                    i = totals[label]
                    if int(str(vc.value).replace(',', '') or 0) != tot_old[i]:
                        self.top_mismatch += 1
                    vc.value = f'{tot_new[i]:,}'
                    changed['top'] += 1


def load_item_dict(secrets_dir):
    p = secrets_dir / 'item_dictionary.json'         # xls와 같은 사전을 공유(같은 회사·품목 체계라는 전제)
    return json.loads(p.read_text(encoding='utf-8')) if p.exists() else {}


def open_for_edit(path):
    wb = openpyxl.load_workbook(str(path))
    return wb, wb.worksheets[0]


def mask_file(masker, src, dst):
    wb, ws = open_for_edit(src)
    changed = masker.run(ws, edit=True)
    wb.save(str(dst))
    return changed
