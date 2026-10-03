import openpyxl
import pandas as pd
import streamlit as st
import xlrd

from masking.core.config import masked_name
from masking.xls.rules import AMOUNT_COLS, HEADER_ROW, TEXT_RULES

from . import _common


def _build_field_labels():
    """열 index -> 원본 표 열 이름(엔진의 실제 이름을 그대로 씀). KEEP_COLS(그대로 유지하는 열)는 마스킹
    대상이 아니므로 넣지 않음. 공급자·공급받는자 양쪽에 '상호'·'대표자명'·'주소'처럼 이름이 겹치는 열이
    있어(홈택스 원본 헤더 자체가 그렇게 두 번 나옴), 겹치는 이름에는 열 index를 붙여 구분한다(겹치지
    않으면 그대로 둠)."""
    raw = {idx: name for idx, (name, _) in TEXT_RULES.items()}
    raw.update(AMOUNT_COLS)
    counts = {}
    for name in raw.values():
        counts[name] = counts.get(name, 0) + 1
    return {idx: (f"{name}(열{idx})" if counts[name] > 1 else name) for idx, name in raw.items()}


FIELD_LABELS = _build_field_labels()


def _describe(path):
    ext = path.suffix.lower().lstrip(".")
    size_bytes = path.stat().st_size
    try:
        if ext == "xls":
            book = xlrd.open_workbook(str(path))
            sheet_count = book.nsheets
            first_sheet_rows = book.sheet_by_index(0).nrows
        else:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_count = len(wb.sheetnames)
            first_sheet_rows = wb[wb.sheetnames[0]].max_row
            wb.close()
        return {
            "파일명": path.name,
            "크기(KB)": round(size_bytes / 1024, 1),
            "확장자": ext,
            "시트 수": sheet_count,
            "첫 시트 행 수": first_sheet_rows,
        }
    except Exception as e:
        return {
            "파일명": path.name,
            "크기(KB)": round(size_bytes / 1024, 1),
            "확장자": ext,
            "시트 수": "읽기 실패",
            "첫 시트 행 수": str(e),
        }


def _read_df(path):
    ext = path.suffix.lower().lstrip(".")
    engine = "xlrd" if ext == "xls" else "openpyxl"
    return pd.read_excel(path, header=HEADER_ROW, engine=engine)


def _head_tail(df, head=3, tail=2):
    return pd.concat([df.head(head), df.tail(tail)])


def _row_with_category(fr):
    return _common.file_result_row(fr, extra={"형식": fr.category})


def _detail_with_category(fr):
    return _common.detail_row(fr, FIELD_LABELS, extra={"형식": fr.category})


def render():
    st.title("Excel Masking")
    st.caption("XLS · XLSX 전자세금계산서 목록 마스킹")

    st.subheader("1. 파일 업로드")
    dirs = _common.upload_section("excel", {".xls", ".xlsx"}, "원본 Excel 파일 업로드 (xls·xlsx, 여러 개 가능)")
    if dirs is None:
        return
    orig_data_dir, masked_data_dir, secrets_dir = dirs

    st.subheader("2. 업로드된 Excel 인식")
    files = sorted(orig_data_dir.glob("*.xls")) + sorted(orig_data_dir.glob("*.xlsx"))

    st.subheader("3. 인식된 Excel 메타데이터")
    if not files:
        st.info("업로드된 파일 중 Excel 파일이 없음.")
        return
    st.dataframe(pd.DataFrame(_describe(p) for p in files), use_container_width=True)

    st.subheader("4. 원본 Excel 미리보기")
    st.caption("마스킹 전 원본 데이터의 앞부분(head)·뒷부분(tail)을 동시에 확인함(아직 마스킹 전이라 실제 값 그대로 보임).")

    file_names = [p.name for p in files]
    sel_name = st.selectbox("미리볼 파일", file_names, key="excel_preview_file")
    sel_path = next(p for p in files if p.name == sel_name)

    try:
        df = _read_df(sel_path)
        st.caption(f"전체 {len(df)}행 (헤더는 {HEADER_ROW + 1}번째 행 기준)")
        st.markdown("**Head(앞 5행)**")
        st.dataframe(df.head(5), use_container_width=True)
        st.markdown("**Tail(뒤 5행)**")
        st.dataframe(df.tail(5), use_container_width=True)
    except Exception as e:
        st.error(f"미리보기 실패: {e}")

    _common.preview_expander(
        "excel_preview", orig_data_dir, masked_data_dir, secrets_dir, {".xls", ".xlsx"},
        caption=(
            "실제로 파일을 만들지 않고, 엔진이 이 Excel들을 어떻게 인식했는지만 확인함(홈택스 전자세금계산서 "
            "목록의 정해진 열 이름을 전제로 하며, 다른 양식이면 여기서 이상하게 보이거나 나중에 \"검토 필요\"로 "
            "걸릴 수 있음). 이 단계에서 `masking_secrets` 폴더가 만들어질 수 있음(비밀 키 생성은 최초 1회만 일어남)."
        ),
    )

    st.subheader("5. 마스킹 실행")
    st.caption(
        f"실제로 마스킹해 `{masked_data_dir}/xls` 또는 `{masked_data_dir}/xlsx`(형식별)에 저장함(폴더가 없으면 "
        "새로 만들고, 이미 있으면 그대로 사용함). 유출 스캔·구조 검증을 통과한 파일만 저장되고, 걸린 파일은 "
        "저장하지 않고 아래 \"검토 필요\"에만 표시됨."
    )
    _common.run_button("excel_run_result", orig_data_dir, masked_data_dir, secrets_dir,
                       {".xls", ".xlsx"}, row_fn=_row_with_category)

    result = st.session_state.get("excel_run_result")
    saved_where = "-"
    if result and result.ok:
        cats = sorted({fr.category for fr in result.ok})
        saved_where = ", ".join(f"`{masked_data_dir / c}`" for c in cats)
    result = _common.render_results(
        "excel_run_result", masked_data_dir, secrets_dir,
        saved_where=saved_where, row_fn=_row_with_category, key_prefix="excel",
    )
    if result and result.ok:
        st.subheader("7. 마스킹 세부내용")
        st.caption("저장된 파일에서 실제로 어떤 항목을 몇 건 바꿨는지 보여줌(항목 이름은 원본 표의 열 이름을 그대로 씀).")
        st.dataframe(pd.DataFrame(_detail_with_category(fr) for fr in result.ok), use_container_width=True)

        st.subheader("8. 마스킹 전후 비교")
        st.caption("저장된 파일 중 하나를 골라 원본과 마스킹 결과의 첫 3행·마지막 2행을 비교함.")
        saved_names = [fr.path.name for fr in result.ok]
        cmp_name = st.selectbox("비교할 파일", saved_names, key="excel_compare_file")
        cmp_fr = next(fr for fr in result.ok if fr.path.name == cmp_name)
        masked_path = masked_data_dir / cmp_fr.category / masked_name(cmp_fr.path.name)

        try:
            st.markdown("**원본**")
            st.dataframe(_head_tail(_read_df(cmp_fr.path)), use_container_width=True)
            st.markdown("**마스킹 결과**")
            st.dataframe(_head_tail(_read_df(masked_path)), use_container_width=True)
        except Exception as e:
            st.error(f"비교 실패: {e}")

    if result is not None:
        st.subheader("9. 결과 다운로드")
        _common.download_buttons(masked_data_dir, secrets_dir, key_prefix="excel")
